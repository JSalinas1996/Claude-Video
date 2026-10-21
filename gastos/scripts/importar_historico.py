"""Convierte las hojas mensuales de Presupuesto_mensual.xlsx (Mayo 26 en adelante)
al formato de la hoja Movimientos. Uso:
  python3 importar_historico.py Presupuesto_mensual.xlsx Gastos.xlsx

Sale en .xlsx (no CSV) para que al subirlo a Drive las fechas y montos queden tipados
y "2026-10" / "06/09" sigan siendo texto, sin depender de la configuración regional.
"""
import datetime as dt
import sys

import openpyxl

MESES = {"Enero": 1, "Febrero": 2, "Marzo": 3, "Abril": 4, "Mayo": 5, "Junio": 6, "Julio": 7,
         "Agosto": 8, "Septiembre": 9, "Octubre": 10, "Noviembre": 11, "Diciembre": 12}
EXCEL_EPOCH = dt.date(1899, 12, 30)

# Gastos fijos -> (categoría, subcategoría). Lo que no esté acá va a Otros.
FIJOS = {
    "Expensas": ("🏠 Hogar", "Expensas"),
    "Alquiler cochera": ("🏠 Hogar", "Cochera"),
    "Luz": ("🏠 Hogar", "Luz"),
    "Gas": ("🏠 Hogar", "Gas"),
    "Internet": ("🏠 Hogar", "Internet"),
    "Telefonía (Tuenti)": ("🏠 Hogar", "Telefonía (Tuenti)"),
    "Obra social": ("❓ Otros", "Salud/Obra Social"),
    "Colegio Balta": ("❓ Otros", "Colegio Balta"),
    "Patente": ("🚗 Transporte", "Patente"),
    "HBO Max": ("⚽ Entretenimiento", "Streaming"),
    "Spotify": ("⚽ Entretenimiento", "Streaming"),
    "YouTube Premium": ("⚽ Entretenimiento", "Streaming"),
    "Disney": ("⚽ Entretenimiento", "Streaming"),
    "Telefonía (Personal)": ("🏠 Hogar", "Telefonía (Personal)"),
    "Ahorro USD": ("💰 Ahorro USD", "Compra USD"),
}

CATEGORIAS = {"🚗 Transporte", "🛒 Supermercado", "⚽ Entretenimiento", "🏠 Hogar", "👕 Ropa/Compras",
              "💳 Impuestos/Financiero", "🍽 Comida/Delivery", "🛡 Seguros", "❓ Otros", "💰 Ahorro USD"}

COLUMNAS = ["ID", "Fecha", "Mes", "Tipo", "Descripción", "Monto", "Categoría", "Clase",
            "Subcategoría", "Medio", "Cuota", "Origen", "Creado"]


def hojas_mensuales(wb):
    for ws in wb:
        partes = ws.title.split()
        if len(partes) == 2 and partes[0] in MESES and partes[1] == "26" and MESES[partes[0]] >= 5:
            yield ws, dt.date(2026, MESES[partes[0]], 1)


def serial_invertido(valor):
    """Un 'dd/mm' que Excel leyó como mm/dd quedó guardado como serial: lo devuelve a (dd, mm, año)."""
    d = EXCEL_EPOCH + dt.timedelta(days=int(valor))
    return d.month, d.day, d.year


def parse_fecha(valor, mes):
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, (int, float)):
        real = EXCEL_EPOCH + dt.timedelta(days=int(valor))
        if real.day > 12:  # no pudo haberse invertido: es la fecha real
            return real
        dia, m, anio = serial_invertido(valor)
        return dt.date(anio, m, dia)
    if isinstance(valor, str) and valor.count("/") == 2:
        d, m, a = valor.split("/")
        return dt.date(int(a), int(m), int(d))
    return mes


def fin_de_mes(mes):
    return (mes.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)


def parse_cuota(valor):
    if isinstance(valor, (int, float)):
        n, total, _ = serial_invertido(valor)
        if n > total:
            n, total = total, n
        return f"{n:02d}/{total:02d}"
    if isinstance(valor, str) and valor.strip("/ "):
        return valor.strip()
    return ""


def filas(ws, mes):
    mes_txt = mes.strftime("%Y-%m")
    rows = list(ws.iter_rows(values_only=True))
    seccion = None
    for r in rows:
        a, b, c, d = (r + (None,) * 9)[:4]
        if a == "💰 INGRESOS":
            seccion = "ingresos"; continue
        if a == "📝 GASTOS FIJOS MENSUALES":
            seccion = "fijos"; continue
        if a == "💳 RESUMEN DE TARJETA":
            seccion = "pre-tarjeta"; continue
        if seccion == "pre-tarjeta" and a == "Fecha" and b == "Establecimiento":
            seccion = "tarjeta"; continue
        if a == "TOTAL CONSUMOS":
            seccion = None; continue

        if seccion == "ingresos" and a and a != "Concepto" and isinstance(c, (int, float)) and c:
            yield dict(Fecha=mes, Mes=mes_txt, Tipo="Ingreso", Descripción=a, Monto=c,
                       Categoría="💵 Ingreso", Clase="", Subcategoría=a, Medio="", Cuota="")
        elif seccion == "fijos" and b and b not in ("Concepto", "TOTAL GASTOS FIJOS") \
                and isinstance(c, (int, float)) and c:
            cat, sub = FIJOS.get(b, ("❓ Otros", b))
            yield dict(Fecha=mes, Mes=mes_txt, Tipo="Gasto", Descripción=b, Monto=c,
                       Categoría=cat, Clase="Fijo", Subcategoría=sub, Medio="Fijo mensual", Cuota="")
        elif seccion == "tarjeta" and b and isinstance(d, (int, float)) and d:
            e, f, g = r[4], r[5], r[6]
            desc = " ".join(str(b).replace("\xa0", " ").split()).rstrip(" $")
            fecha = parse_fecha(a, mes)
            if fecha > fin_de_mes(mes) and fecha.day <= 12:  # día y mes cargados al revés
                fecha = dt.date(fecha.year, fecha.day, fecha.month)
            cat, sub = e or "❓ Otros", g or ""
            if cat not in CATEGORIAS:
                cat, sub = "❓ Otros", sub or cat
            yield dict(Fecha=fecha, Mes=mes_txt, Tipo="Gasto", Descripción=desc, Monto=d,
                       Categoría=cat, Clase=f or "Variable", Subcategoría=sub,
                       Medio="Tarjeta de crédito", Cuota=parse_cuota(c))


def main(origen, destino):
    wb = openpyxl.load_workbook(origen, data_only=True)
    creado = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    salida = openpyxl.Workbook()
    ws_out = salida.active
    ws_out.title = "Movimientos"
    ws_out.append(COLUMNAS)
    n = 0
    for ws, mes in sorted(hojas_mensuales(wb), key=lambda x: x[1]):
        for i, fila in enumerate(filas(ws, mes)):
            fila.update(ID=f"hist-{mes:%Y%m}-{i:03d}", Origen="histórico", Creado=creado,
                        Monto=round(fila["Monto"], 2))
            ws_out.append([fila[c] for c in COLUMNAS])
            n += 1
    for celda in ws_out["B"][1:]:
        celda.number_format = "yyyy-mm-dd"
    for col in ("C", "K"):
        for celda in ws_out[col][1:]:
            celda.number_format = "@"
    for celda in ws_out["F"][1:]:
        celda.number_format = "#,##0.00"
    ws_out.freeze_panes = "A2"
    salida.save(destino)
    print(f"{n} filas -> {destino}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
